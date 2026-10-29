"""Turn a JD link or uploaded file (docx/pdf/xlsx/csv/txt) into plain text."""
import csv
import html
import io
import json
import re
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup

UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"


def _clean(text: str) -> str:
    text = re.sub(r"[ \t\xa0]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _html_text(fragment: str) -> str:
    return _clean(BeautifulSoup(fragment or "", "html.parser").get_text("\n"))


def _jd(title: str, company: str, location: str, body: str, host: str) -> dict:
    head = "\n".join(x for x in (title, f"{company} — {location}" if location else company) if x)
    return {"text": _clean(f"{head}\n\n{body}")[:20000], "page_title": title, "title": title,
            "company": company, "location": location, "site_name": company, "host": host}


def _slug_name(slug: str) -> str:
    return " ".join(w.capitalize() for w in re.split(r"[-_]", slug) if w)


async def _ats_posting(c: httpx.AsyncClient, url: str) -> dict | None:
    """Job boards that render client-side (Ashby, Greenhouse, Lever) expose public posting APIs."""
    u = urlparse(url)
    host, parts = (u.hostname or "").lower(), [p for p in u.path.split("/") if p]
    if host == "jobs.ashbyhq.com" and len(parts) >= 2:
        org, jid = parts[0], parts[1]
        r = await c.get(f"https://api.ashbyhq.com/posting-api/job-board/{org}", params={"includeCompensation": "true"})
        r.raise_for_status()
        job = next((j for j in r.json().get("jobs", []) if j.get("id") == jid), None)
        if job:
            body = job.get("descriptionPlain") or _html_text(job.get("descriptionHtml", ""))
            meta = " | ".join(x for x in (job.get("department"), job.get("team"), job.get("employmentType")) if x)
            return _jd(job.get("title", ""), _slug_name(org), job.get("location", ""), f"{meta}\n\n{body}", host)
    if host.endswith("greenhouse.io") and "jobs" in parts:
        i = parts.index("jobs")
        org, jid = (parts[i - 1] if i else ""), (parts[i + 1] if len(parts) > i + 1 else "")
        if org and jid.isdigit():
            r = await c.get(f"https://boards-api.greenhouse.io/v1/boards/{org}/jobs/{jid}")
            r.raise_for_status()
            job = r.json()
            b = await c.get(f"https://boards-api.greenhouse.io/v1/boards/{org}")
            company = b.json().get("name") if b.status_code == 200 else _slug_name(org)
            return _jd(job.get("title", ""), company, (job.get("location") or {}).get("name", ""),
                       _html_text(html.unescape(job.get("content", ""))), host)
    if host == "jobs.lever.co" and len(parts) >= 2:
        org, jid = parts[0], parts[1]
        r = await c.get(f"https://api.lever.co/v0/postings/{org}/{jid}")
        r.raise_for_status()
        job = r.json()
        lists = "\n\n".join(f"{x.get('text', '')}\n{_html_text(x.get('content', ''))}" for x in job.get("lists", []))
        body = "\n\n".join(x for x in (job.get("descriptionPlain"), lists, job.get("additionalPlain")) if x)
        return _jd(job.get("text", ""), _slug_name(org), (job.get("categories") or {}).get("location", ""), body, host)
    return None


def _json_ld_posting(soup: BeautifulSoup, host: str) -> dict | None:
    """schema.org JobPosting embedded by many career sites."""
    for tag in soup.find_all("script", type="application/ld+json"):
        try:
            data = json.loads(tag.string or "")
        except (ValueError, TypeError):
            continue
        items = data if isinstance(data, list) else data.get("@graph", [data]) if isinstance(data, dict) else []
        for it in items:
            if isinstance(it, dict) and it.get("@type") == "JobPosting" and len(it.get("description", "")) > 200:
                org = it.get("hiringOrganization") or {}
                loc = it.get("jobLocation") or {}
                loc = loc[0] if isinstance(loc, list) and loc else loc
                addr = loc.get("address", {}) if isinstance(loc, dict) else {}
                addr = addr if isinstance(addr, dict) else {}
                location = ", ".join(x for x in (addr.get("addressLocality"), addr.get("addressCountry")) if isinstance(x, str) and x)
                return _jd(it.get("title", ""), org.get("name", "") if isinstance(org, dict) else "", location,
                           _html_text(html.unescape(it["description"])), host)
    return None


async def fetch_url(url: str) -> dict:
    async with httpx.AsyncClient(follow_redirects=True, timeout=20, headers={"User-Agent": UA}) as c:
        try:
            posting = await _ats_posting(c, url)
        except (httpx.HTTPError, ValueError):
            posting = None
        if posting and len(posting["text"]) > 200:
            return posting
        r = await c.get(url)
        r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    host = urlparse(url).hostname or ""
    posting = _json_ld_posting(soup, host)
    if posting:
        return posting
    title = (soup.title.string or "").strip() if soup.title else ""
    og_site = soup.find("meta", property="og:site_name")
    for tag in soup(["script", "style", "noscript", "svg", "header", "footer", "nav", "form", "iframe"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    text = main.get_text("\n")
    return {
        "text": _clean(text)[:20000],
        "page_title": title,
        "site_name": og_site["content"].strip() if og_site and og_site.get("content") else "",
        "host": host,
    }


def parse_docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return _clean("\n".join(parts))


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return _clean("\n".join((p.extract_text() or "") for p in reader.pages))


def parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        # Key/value sheet (2 columns) or a table with headers: render as "Header: value" lines.
        for row in rows[1:] if any(header) and len(rows) > 1 else rows:
            cells = ["" if v is None else str(v).strip() for v in row]
            if not any(cells):
                continue
            if any(header) and len(rows) > 1 and len(header) > 2:
                lines.extend(f"{h}: {v}" for h, v in zip(header, cells) if v)
                lines.append("")
            else:
                lines.append(" : ".join(c for c in cells if c))
        if any(header) and len(header) <= 2:
            lines.insert(0, " : ".join(h for h in header if h))
    return _clean("\n".join(lines))


def parse_csv(data: bytes) -> str:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8", "ignore"))))
    return _clean("\n".join(" : ".join(c for c in r if c) for r in rows))


def parse_file(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        return parse_docx(data)
    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    if name.endswith(".csv"):
        return parse_csv(data)
    if name.endswith((".txt", ".md")):
        return _clean(data.decode("utf-8", "ignore"))
    raise ValueError("Unsupported file type. Use .docx, .pdf, .xlsx, .csv or .txt")
